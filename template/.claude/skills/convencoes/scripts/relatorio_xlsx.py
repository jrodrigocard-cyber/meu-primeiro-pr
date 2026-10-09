#!/usr/bin/env python3
"""Gera Excel (padrão Anályse) a partir do mesmo JSON do relatório PDF.
Uso: relatorio_xlsx.py resumo.json [saida.xlsx]  (padrão: ~/Downloads/<nome>.xlsx)"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "008BD0"
LOGO = Path(__file__).resolve().parents[4] / "assets" / "logo-analyse.png"
fino = Border(bottom=Side(style="thin", color="D9D9D9"))
quebra = Alignment(wrap_text=True, vertical="top")


def main(js, saida):
    d = json.load(open(js, encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", (34, 100, 22)):
        ws.column_dimensions[col].width = w
    if LOGO.exists():
        img = Image(str(LOGO))
        img.height, img.width = 62, 214
        ws.add_image(img, "A1")
    ws.row_dimensions[1].height = 50
    ws["A3"] = d["titulo"]
    ws["A3"].font = Font(name="Arial", size=15, bold=True)
    r = 5
    for k, v in d["meta"]:
        ws.cell(r, 1, k.upper()).font = Font(name="Arial", size=9, bold=True, color=AZUL)
        c = ws.cell(r, 2, v)
        c.font = Font(name="Arial", size=10)
        c.alignment = quebra
        ws.cell(r, 1).alignment = quebra
        r += 1
    r += 1
    for i, sec in enumerate(d["secoes"], 1):
        ws.cell(r, 1, f"{i}. {sec['titulo']}").font = Font(name="Arial", size=12, bold=True)
        r += 1
        if "tabela" in sec:
            for j, h in enumerate(("Tema", "O que diz", "Cláusula"), 1):
                c = ws.cell(r, j, h)
                c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=AZUL)
                c.alignment = Alignment(vertical="center")
            r += 1
            for t, txt, cl in sec["tabela"]:
                for j, v in enumerate((t, txt, cl), 1):
                    c = ws.cell(r, j, v)
                    c.font = Font(name="Arial", size=10, bold=(j == 1))
                    c.alignment = quebra
                    c.border = fino
                ws.row_dimensions[r].height = max(18, 14 * (len(txt) // 105 + 1))
                r += 1
        else:
            for b in sec["itens"]:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
                c = ws.cell(r, 1, "• " + b)
                c.font = Font(name="Arial", size=10)
                c.alignment = quebra
                ws.row_dimensions[r].height = max(18, 14 * (len(b) // 150 + 1))
                r += 1
        r += 1
    ws.cell(r, 1, f"Anályse Assessoria Contábil S/S - {d['titulo']} - {d['data']}").font = Font(name="Arial", size=8, color="8C8C8C")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(saida)


if __name__ == "__main__":
    js = sys.argv[1]
    if len(sys.argv) > 2:
        out = sys.argv[2]
    else:
        pasta = Path.home() / "Downloads"
        pasta.mkdir(exist_ok=True)
        out = str(pasta / (Path(js).stem + ".xlsx"))
    main(js, out)
    print("Excel salvo em", out)
